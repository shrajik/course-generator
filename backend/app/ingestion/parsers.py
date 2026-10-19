"""Turn an uploaded file into an ordered list of headings, paragraphs and tables.

Three formats: PDF (text layer only - a scanned PDF is refused with a clear
message rather than silently producing nothing), DOCX (real heading styles and
tables) and plain text / Markdown. The parsers are pure functions of bytes, so
they are cheap to test and are run off the event loop by the caller.

Uploaded files are untrusted: sizes are bounded, a DOCX is checked for an
oversized zip before it is opened, and nothing here executes or fetches
anything from the file.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Literal

MAX_PAGES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
# A text layer thinner than this per page means the PDF is scanned images.
MIN_CHARS_PER_PAGE = 20


class IngestionError(Exception):
    """A file that cannot be turned into text. The message is shown to the user."""


@dataclass
class ParsedBlock:
    kind: Literal["heading", "paragraph", "table"]
    text: str
    level: int = 0  # headings only: 1 = top
    page: int | None = None


@dataclass
class ParsedDocument:
    blocks: list[ParsedBlock] = field(default_factory=list)
    page_count: int | None = None

    @property
    def word_count(self) -> int:
        return sum(len(block.text.split()) for block in self.blocks)

    def outline(self, limit: int = 40) -> list[str]:
        return [
            ("  " * (block.level - 1)) + block.text
            for block in self.blocks
            if block.kind == "heading"
        ][:limit]


# --- shared helpers ------------------------------------------------------------

_NUMBERED_HEADING_RE = re.compile(r"^\d+(?:\.\d+){0,4}[.)]?\s+\S")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\x00", "")).strip()


def looks_like_heading(line: str) -> bool:
    """A short line that is not a sentence: ALL CAPS, or a numbered title."""
    text = line.strip()
    if not text or len(text) > 80 or len(text.split()) > 10:
        return False
    if text[-1] in ".,;:!?":
        return False
    if text.isupper() and any(char.isalpha() for char in text) and len(text) > 2:
        return True
    return bool(_NUMBERED_HEADING_RE.match(text)) and not text.replace(".", "").isdigit()


def _heading_level(text: str) -> int:
    match = re.match(r"^(\d+(?:\.\d+){0,4})", text)
    return min(match.group(1).count(".") + 1, 4) if match else 1


def _decode(data: bytes) -> str:
    # UTF-16 is only trusted with a byte-order mark: without one, arbitrary
    # 8-bit text happens to "decode" as UTF-16 into nonsense.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeError:
        return data.decode("cp1252", errors="replace")


# --- plain text / markdown -----------------------------------------------------


def parse_text(data: bytes, *, markdown: bool) -> ParsedDocument:
    text = _decode(data).replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise IngestionError("The file is empty.")
    blocks: list[ParsedBlock] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(ParsedBlock("paragraph", _clean(" ".join(paragraph))))
            paragraph.clear()

    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        match = _MD_HEADING_RE.match(line) if markdown else None
        if match:
            flush()
            blocks.append(ParsedBlock("heading", _clean(match.group(2)), level=len(match.group(1))))
        elif not paragraph and looks_like_heading(line):
            blocks.append(ParsedBlock("heading", _clean(line), level=_heading_level(line)))
        else:
            paragraph.append(line.strip())
    flush()
    return ParsedDocument(blocks=blocks)


# --- docx ----------------------------------------------------------------------


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise IngestionError("The document expands to an unreasonable size.")
            if "word/document.xml" not in archive.namelist():
                raise IngestionError("This is not a Word (.docx) document.")
    except zipfile.BadZipFile as exc:
        raise IngestionError("This is not a valid Word (.docx) document.") from exc


def parse_docx(data: bytes) -> ParsedDocument:
    _check_zip(data)
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - any corrupt-file failure reads the same to the user
        raise IngestionError("The Word document could not be opened.") from exc

    blocks: list[ParsedBlock] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = _clean(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            heading = re.match(r"(?i)^heading\s*(\d)", style)
            if heading:
                blocks.append(ParsedBlock("heading", text, level=min(int(heading.group(1)), 6)))
            elif style.lower() == "title":
                blocks.append(ParsedBlock("heading", text, level=1))
            elif style.lower().startswith("list"):
                blocks.append(ParsedBlock("paragraph", "- " + text))
            else:
                blocks.append(ParsedBlock("paragraph", text))
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = _clean(cell.text)
                    if not cells or cells[-1] != value:  # merged cells repeat
                        cells.append(value)
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(ParsedBlock("table", "\n".join(rows)))
    if not blocks:
        raise IngestionError("The Word document has no text.")
    return ParsedDocument(blocks=blocks)


# --- pdf -----------------------------------------------------------------------


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise IngestionError("The PDF is password-protected.")
        pages = len(reader.pages)
    except IngestionError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise IngestionError("The PDF could not be read.") from exc
    except Exception as exc:  # noqa: BLE001
        raise IngestionError("The PDF could not be read.") from exc

    if pages > MAX_PAGES:
        raise IngestionError(f"The PDF has {pages} pages; the limit is {MAX_PAGES}.")

    blocks: list[ParsedBlock] = []
    characters = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - one bad page must not lose the rest
            text = ""
        characters += len(text.strip())
        paragraph: list[str] = []

        def flush(number: int = number, paragraph: list[str] = paragraph) -> None:
            if paragraph:
                blocks.append(ParsedBlock("paragraph", _clean(" ".join(paragraph)), page=number))
                paragraph.clear()

        for line in text.replace("\r", "\n").split("\n"):
            if not line.strip():
                flush()
            elif looks_like_heading(line):
                flush()
                blocks.append(
                    ParsedBlock("heading", _clean(line), level=_heading_level(line), page=number)
                )
            else:
                paragraph.append(line.strip())
        flush()

    if characters < MIN_CHARS_PER_PAGE * max(pages, 1):
        raise IngestionError(
            "This PDF has no readable text - it looks like scanned images. "
            "Scanned PDFs need OCR, which is not supported yet; upload a text version instead."
        )
    return ParsedDocument(blocks=blocks, page_count=pages)


def parse(kind: str, data: bytes) -> ParsedDocument:
    if kind == "pdf":
        return parse_pdf(data)
    if kind == "docx":
        return parse_docx(data)
    if kind in {"txt", "md"}:
        return parse_text(data, markdown=kind == "md")
    raise IngestionError(f"Unsupported file type: {kind}")
