"""Uploaded DOCX template -> registry -> generation -> template-aware PDF.

Covers the 16 checkpoints in the feature brief. The DOCX fixture is built
in-memory with python-docx rather than committed as a binary, so what each
test asserts about the parse is visible next to the document that produced it.

The load-bearing assertions are:
* `test_builtin_*` - technical_v1/non_technical_v1 are byte-identical to
  before, because the whole feature is additive.
* `test_end_to_end_*` - an uploaded template reaches the rendered HTML, which
  is the exact step the audit found missing.
"""

from __future__ import annotations

import io

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.course.templates.docx_parser import (
    PARSER_VERSION,
    context_from,
    names_in,
    parse_docx_template,
    resolve,
    unknown_names,
)
from app.course.templates.registry import (
    UPLOADED_PREFIX,
    is_known_template_id,
    is_uploaded_template_id,
    load_template,
    register_uploaded_template,
)
from app.render.html_renderer import render_document_html
from app.schemas.blocks import BlockType
from app.schemas.document import Block, BlockMeta, CourseDocument, DocumentMeta, Page, PageSize
from app.schemas.template import CourseTemplate, PageGeometry, TemplateTheme

BRAND = RGBColor(0x1F, 0x4E, 0x79)


def _fixture_docx() -> bytes:
    """A template exercising everything the parser claims to read."""
    document = docx.Document()

    section = document.sections[0]
    section.page_width, section.page_height = 7772400, 10058400   # 8.5in x 11in
    section.left_margin = section.right_margin = 685800           # 0.75in
    section.top_margin = section.bottom_margin = 457200           # 0.5in
    section.header.paragraphs[0].text = "{{COURSE_TITLE}} | {{AUDIENCE}}"
    section.footer.paragraphs[0].text = "Acme Corp Confidential"

    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    heading = document.styles["Heading 1"]
    heading.font.name = "Verdana"
    heading.font.color.rgb = BRAND

    document.add_paragraph("Acme Handbook", style="Title")

    document.add_paragraph("Introduction", style="Heading 1")
    document.add_paragraph("Sets up {{COURSE_TITLE}} for {{AUDIENCE}}.")
    document.add_paragraph("Why This Matters", style="Heading 2")

    document.add_paragraph("Learning Objectives (required)", style="Heading 1")
    document.add_paragraph("{{OBJECTIVES}}")

    document.add_paragraph("Knowledge Check", style="Heading 1")
    document.add_paragraph("Multiple Choice", style="Heading 2")

    document.add_paragraph("Glossary (optional)", style="Heading 1")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "{{UNSUPPORTED_THING}}"

    document.add_paragraph("Version History", style="Heading 1")
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def docx_bytes() -> bytes:
    return _fixture_docx()


@pytest.fixture(scope="module")
def parsed(docx_bytes: bytes):
    return parse_docx_template(
        docx_bytes, template_id="uploaded:test-0001", name="Acme Handbook", kind="non_technical"
    )


# --- 2/3. parse + structure extraction ---------------------------------------


def test_parse_produces_a_valid_course_template(parsed) -> None:
    """The parser's output is an ordinary CourseTemplate, which is what lets
    every downstream consumer stay unchanged."""
    assert isinstance(parsed.template, CourseTemplate)
    CourseTemplate.model_validate(parsed.template.model_dump(mode="json"))
    assert parsed.template.template_id == "uploaded:test-0001"
    assert parsed.template.kind == "non_technical"
    assert parsed.parser_version == PARSER_VERSION


def test_sections_are_extracted_in_document_order(parsed) -> None:
    assert [s.key for s in parsed.template.sections] == [
        "introduction",
        "learning_objectives",
        "knowledge_check",
        "glossary",
        "version_history",
    ]


def test_explicit_required_marker_is_honoured_and_optional_is_not(parsed) -> None:
    """Obligation is only set where the document says so - the brief forbids
    inventing rules the DOCX does not express."""
    by_key = {s.key: s for s in parsed.template.sections}
    assert by_key["learning_objectives"].required is True
    assert by_key["glossary"].required is False
    assert by_key["introduction"].required is False
    assert parsed.structure.has_explicit_requirements is True


def test_obligation_markers_are_stripped_from_labels(parsed) -> None:
    by_key = {s.key: s.label for s in parsed.template.sections}
    assert by_key["learning_objectives"] == "Learning Objectives"
    assert by_key["glossary"] == "Glossary"


def test_block_types_are_inferred_only_from_known_keywords(parsed) -> None:
    by_key = {s.key: s.block_types for s in parsed.template.sections}
    assert BlockType.QUIZ in by_key["knowledge_check"]
    assert BlockType.LEARNING_OBJECTIVES in by_key["learning_objectives"]
    assert BlockType.TABLE in by_key["glossary"]
    # "Introduction" matches no keyword, so nothing is guessed for it.
    assert by_key["introduction"] == []


def test_subheadings_become_section_guidance(parsed) -> None:
    intro = next(s for s in parsed.template.sections if s.key == "introduction")
    assert "Why This Matters" in intro.guidance


def test_history_does_not_match_the_story_keyword(parsed) -> None:
    """Regression: substring matching made "Version History" a story block."""
    history = next(s for s in parsed.template.sections if s.key == "version_history")
    assert BlockType.STORY not in history.block_types


def test_nothing_is_forced_on_the_reviewer(parsed) -> None:
    """A DOCX cannot express required block types, so none are declared -
    otherwise the reviewer would block approval on an invented rule."""
    assert parsed.template.required_block_types == []
    assert parsed.template.allowed_block_types == []


# --- 4. style extraction ------------------------------------------------------


def test_typography_is_extracted_from_the_docx(parsed) -> None:
    theme = parsed.template.theme
    assert theme.font_family.startswith("Georgia")
    assert theme.heading_font_family.startswith("Verdana")
    assert theme.base_font_size == 12.0
    assert theme.base_line_height == 1.5


def test_colors_are_extracted_from_the_docx(parsed) -> None:
    assert parsed.template.theme.heading_color == "#1f4e79"
    assert parsed.template.theme.accent_color == "#1f4e79"


def test_page_geometry_is_extracted(parsed) -> None:
    geometry = parsed.template.theme.geometry()
    assert (geometry.width, geometry.height) == (816.0, 1056.0)   # 8.5x11in at 96dpi
    assert geometry.margin_x == 72.0                              # 0.75in
    assert geometry.margin_top == 48.0                            # 0.5in


def test_running_header_and_footer_are_extracted(parsed) -> None:
    assert parsed.template.theme.header_text == "{{COURSE_TITLE}} | {{AUDIENCE}}"
    assert parsed.template.theme.footer_text == "Acme Corp Confidential"


def test_font_names_are_emitted_unquoted(parsed) -> None:
    """Quotes in a font stack are HTML-escaped by Jinja inside <style>, which
    silently breaks the declaration - so the stack must never contain them."""
    assert '"' not in parsed.template.theme.font_family
    assert "'" not in parsed.template.theme.heading_font_family


# --- 7. placeholders ----------------------------------------------------------


def test_supported_and_unsupported_placeholders_are_separated(parsed) -> None:
    report = parsed.report()
    assert "COURSE_TITLE" in report["supported_placeholders"]
    assert "AUDIENCE" in report["supported_placeholders"]
    # Surfaced, not silently ignored.
    assert "UNSUPPORTED_THING" in report["unsupported_placeholders"]


def test_resolve_substitutes_known_values_only() -> None:
    text = "{{COURSE_TITLE}} for {{AUDIENCE}} / {{MYSTERY}}"
    out = resolve(text, {"course_title": "Risk 101", "audience": "Analysts"})
    assert out == "Risk 101 for Analysts / {{MYSTERY}}"


def test_resolve_leaves_a_known_placeholder_with_no_value_visible() -> None:
    """An empty substitution would be an invisible failure; a visible
    {{PLACEHOLDER}} is a bug someone can actually report."""
    assert resolve("{{COURSE_TITLE}}", {}) == "{{COURSE_TITLE}}"


def test_placeholder_helpers() -> None:
    assert names_in("{{A_B}} x {{A_B}} {{C}}") == ["A_B", "C"]
    assert unknown_names(["COURSE_TITLE", "NOPE"]) == ["NOPE"]


def test_context_from_reads_the_learner_profile() -> None:
    from app.schemas.course import CourseInput
    from app.schemas.learner import LearnerProfile

    course_input = CourseInput(
        course_title="Risk 101",
        toc=[{"title": "One"}],
        target_audience="Analysts",
        template="non_technical",
        learner_profile=LearnerProfile(chunk_word_cap=275),
    )
    context = context_from(course_input)
    assert context["course_title"] == "Risk 101"
    assert context["chunk_word_cap"] == 275


# --- 8. registry --------------------------------------------------------------


def test_registry_resolves_an_uploaded_template(parsed) -> None:
    register_uploaded_template(parsed.template)
    loaded = load_template("uploaded:test-0001")
    assert loaded.template_id == "uploaded:test-0001"
    assert len(loaded.sections) == 5
    assert is_known_template_id("uploaded:test-0001")
    assert is_uploaded_template_id("uploaded:test-0001")
    assert not is_uploaded_template_id("technical_v1")


def test_unknown_uploaded_id_reports_a_useful_error() -> None:
    from app.core.errors import NotFoundError

    with pytest.raises(NotFoundError, match="no parsed configuration"):
        load_template(f"{UPLOADED_PREFIX}does-not-exist")


# --- 15/16. built-in templates are untouched ----------------------------------


@pytest.mark.parametrize("template_id", ["technical_v1", "non_technical_v1"])
def test_builtin_templates_declare_no_custom_geometry(template_id: str) -> None:
    """The whole feature is additive: a built-in template must still use the
    layout engine's own page box, or every existing course would re-paginate."""
    template = load_template(template_id)
    assert template.theme.page is None
    assert template.theme.geometry() == PageGeometry()


@pytest.mark.parametrize("template_id", ["technical_v1", "non_technical_v1"])
def test_builtin_theme_extension_fields_are_inert(template_id: str) -> None:
    theme = load_template(template_id).theme
    assert theme.heading_font_family == ""
    assert theme.header_text == ""
    assert theme.footer_text == ""
    assert theme.heading_family() == theme.font_family
    assert theme.heading_ink() == theme.text_color


def test_builtin_kind_resolution_is_unchanged() -> None:
    assert load_template("technical").template_id == "technical_v1"
    assert load_template("non_technical").template_id == "non_technical_v1"


# --- 12/13/14. document + renderer -------------------------------------------


def _document(template: CourseTemplate) -> CourseDocument:
    geometry = template.theme.geometry()
    return CourseDocument(
        document_id="doc_test",
        course_id="crs_test",
        course_title="Risk Basics",
        template_id=template.template_id,
        meta=DocumentMeta(
            audience="Analysts",
            template_name=template.name,
            theme=template.theme.model_dump(mode="json"),
        ),
        pages=[
            Page(
                id="page_1",
                page_number=1,
                kind="content",
                size=PageSize(width=geometry.width, height=geometry.height),
                blocks=[
                    Block(
                        id="block_aaa001",
                        type=BlockType.PARAGRAPH,
                        content={"text": "Body copy."},
                        meta=BlockMeta(),
                    )
                ],
            )
        ],
    )


def test_end_to_end_uploaded_template_styles_the_rendered_html(parsed) -> None:
    """The step the audit found missing: uploaded styling reaching the page."""
    register_uploaded_template(parsed.template)
    template = load_template("uploaded:test-0001")
    html = render_document_html(_document(template), template)

    assert "@page { size: 816.0px 1056.0px" in html   # the DOCX's page box
    assert "Georgia," in html                          # body face
    assert "Verdana," in html                          # heading face
    assert "#1f4e79" in html                           # brand heading colour
    assert "Acme Corp Confidential" in html            # DOCX footer
    # Header placeholders resolved against the document.
    assert "Risk Basics | Analysts" in html
    # Escaped quotes would silently break the font declaration.
    assert "&#34" not in html


def test_end_to_end_builtin_rendering_is_unchanged() -> None:
    template = load_template("technical_v1")
    html = render_document_html(_document(template), template)
    assert "@page { size: 794.0px 1123.0px" in html
    assert "Inter, Segoe UI" in html
    # No running header is introduced for a template that declares none.
    assert 'class="runhead"' not in html


def test_document_records_the_selected_template_id(parsed) -> None:
    register_uploaded_template(parsed.template)
    document = _document(load_template("uploaded:test-0001"))
    assert document.template_id == "uploaded:test-0001"


# --- layout geometry ----------------------------------------------------------


def test_flow_blocks_honours_a_custom_page_box() -> None:
    """Positioning must follow the template's box, or the renderer's @page and
    the block coordinates would disagree and content would clip."""
    from app.course.document.layout import flow_blocks

    def block() -> Block:
        return Block(
            id="block_bbb001",
            type=BlockType.PARAGRAPH,
            content={"text": "word " * 40},
            meta=BlockMeta(),
        )

    narrow = PageGeometry(width=500.0, height=700.0, margin_x=20.0, margin_top=30.0)
    pages = flow_blocks([block()], geometry=narrow)
    placed = pages[0][0]
    assert placed.layout.x == 20.0
    assert placed.layout.y == 30.0
    assert placed.layout.width == narrow.content_width

    default_pages = flow_blocks([block()])
    assert default_pages[0][0].layout.x == PageGeometry().margin_x


# --- 10. validation report ----------------------------------------------------


def test_parse_report_shape(parsed) -> None:
    report = parsed.report()
    assert report["parser_version"] == PARSER_VERSION
    assert report["detected"]["sections"] == 5
    assert report["detected"]["tables"] == 1
    assert report["detected"]["explicit_required_markers"] is True
    assert isinstance(report["warnings"], list)
    assert {"key", "label", "required", "block_types"} <= set(report["sections"][0])


def test_a_docx_with_no_headings_still_parses_and_warns() -> None:
    """A style-only template is valid: the brief says not to reject a document
    merely because part of it cannot be used."""
    document = docx.Document()
    document.add_paragraph("Just prose, no headings at all.")
    buffer = io.BytesIO()
    document.save(buffer)

    result = parse_docx_template(
        buffer.getvalue(), template_id="uploaded:bare", name="Bare", kind="technical"
    )
    assert result.template.sections == []
    assert any("No Heading 1" in warning for warning in result.warnings)


def test_unreadable_file_raises_for_the_caller_to_turn_into_a_validation_error() -> None:
    with pytest.raises(Exception):
        parse_docx_template(b"not a docx at all", template_id="uploaded:x", name="X", kind="technical")


def test_theme_defaults_match_the_layout_engine() -> None:
    """If these drift apart, every built-in course silently re-paginates."""
    from app.schemas.document import PAGE_HEIGHT, PAGE_MARGIN_X, PAGE_WIDTH

    geometry = TemplateTheme().geometry()
    assert (geometry.width, geometry.height) == (PAGE_WIDTH, PAGE_HEIGHT)
    assert geometry.margin_x == PAGE_MARGIN_X


# --- 1/6/9/10/11/12/13/14. full pipeline, offline ------------------------------


def test_uploaded_template_drives_a_whole_generated_course(client, parsed) -> None:
    """The complete chain the brief asks for, end to end and database-free:

    uploaded DOCX -> registry -> CourseInput.template_id_override -> planner
    -> writer -> reviewer -> CourseDocument -> template-aware renderer.

    Runs against the offline mock AI client (tests/conftest.py sets
    MOCK_OPENAI=true), so it asserts that the *template* propagates, not that
    a particular model writes particular prose.
    """
    register_uploaded_template(parsed.template)

    created = client.post(
        "/api/courses",
        json={
            "course_title": "Risk Basics",
            "toc": [{"title": "What Risk Means"}],
            "target_audience": "Analysts",
            "template": "non_technical",
            "template_id_override": "uploaded:test-0001",
        },
    )
    assert created.status_code == 201, created.text
    course_id = created.json()["course_id"]
    document_id = created.json()["document_id"]

    # The selection survives persistence rather than being dropped on the way
    # in, which is exactly what the audit found happening before.
    fetched = client.get(f"/api/courses/{course_id}").json()["course"]
    assert fetched["input"]["template_id_override"] == "uploaded:test-0001"
    assert fetched["template_id"] == "uploaded:test-0001"

    # The planner resolved the uploaded template, not technical_v1.
    blueprint = client.get(f"/api/courses/{course_id}/blueprint").json()
    assert blueprint["template_id"] == "uploaded:test-0001"

    result = client.post(f"/api/courses/{course_id}/generate", json={"mode": "sync"}).json()
    assert result["status"] == "ready", result

    document = client.get(f"/api/documents/{document_id}").json()
    assert document["template_id"] == "uploaded:test-0001"
    # The document carries the DOCX's own typography, which is what the
    # renderer reads when it builds the stylesheet.
    assert document["meta"]["theme"]["font_family"].startswith("Georgia")
    assert document["pages"][0]["size"]["width"] == 816.0

    # Render through the real PDF path's HTML step (Playwright itself needs a
    # browser, which the offline suite does not assume).
    from app.services.pdf_service import PdfService

    rendered = PdfService().render_html(CourseDocument.model_validate(document))
    assert "@page { size: 816.0px 1056.0px" in rendered
    assert "Georgia," in rendered
    assert "Acme Corp Confidential" in rendered


def test_builtin_course_is_unaffected_by_the_feature(client) -> None:
    """Same request without an override still resolves to technical_v1 and
    still renders at the built-in page size."""
    created = client.post(
        "/api/courses",
        json={
            "course_title": "Plain Course",
            "toc": [{"title": "Chapter One"}],
            "target_audience": "Engineers",
            "template": "technical",
        },
    )
    assert created.status_code == 201, created.text
    course_id = created.json()["course_id"]
    fetched = client.get(f"/api/courses/{course_id}").json()["course"]
    assert fetched["template_id"] == "technical_v1"
    assert fetched["input"]["template_id_override"] is None


def test_unknown_template_override_is_rejected(client) -> None:
    response = client.post(
        "/api/courses",
        json={
            "course_title": "Bad Template",
            "toc": [{"title": "One"}],
            "target_audience": "Anyone",
            "template": "technical",
            "template_id_override": "uploaded:nope",
        },
    )
    assert response.status_code == 422, response.text
    assert "Unknown template_id_override" in response.text
